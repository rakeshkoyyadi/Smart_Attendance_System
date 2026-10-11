import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from database import get_connection


def generate_excel_report():

    conn = get_connection()

    if conn is None:
        return

    cursor = conn.cursor()

    # Get distinct attendance dates
    cursor.execute("""
        SELECT DISTINCT attendance_date
        FROM attendance
        ORDER BY attendance_date
    """)

    dates = cursor.fetchall()

    print("\n========== ATTENDANCE DATES ==========")
    if not dates:
        print("No attendance dates found!")
    else:
        for d in dates:
            print(d[0])
    print("======================================\n")

    os.makedirs("reports", exist_ok=True)

    filename = os.path.abspath("reports/Attendance_Report.xlsx")
    print("Saving Excel to:", filename)

    # Open existing workbook or create new
    if os.path.exists(filename):
        workbook = load_workbook(filename)
        print("Existing workbook opened.")
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)
        print("New workbook created.")

    for row in dates:

        attendance_date = row[0]
        sheet_name = str(attendance_date)

        print(f"\nProcessing Sheet: {sheet_name}")

        # Delete old sheet
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]

        sheet = workbook.create_sheet(title=sheet_name)

        headers = [
            "Student ID",
            "Student Name",
            "Attendance Date",
            "Entry Time",
            "Exit Time",
            "Status"
        ]

        sheet.append(headers)

        for cell in sheet[1]:
            cell.font = Font(bold=True)

        cursor.execute("""
            SELECT
                student_id,
                student_name,
                attendance_date,
                entry_time,
                exit_time,
                status
            FROM attendance
            WHERE attendance_date = %s
            ORDER BY student_id
        """, (attendance_date,))

        records = cursor.fetchall()

        print(f"Records Found: {len(records)}")

        # Print first 5 records
        for record in records[:5]:
            print(record)

        for record in records:
            sheet.append(record)

        # Auto-adjust column widths
        for column in sheet.columns:

            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))

            sheet.column_dimensions[column_letter].width = max_length + 3

    workbook.save(filename)

    print("\nWorkbook Sheets:")
    print(workbook.sheetnames)

    print("\n========================================")
    print("Excel Report Generated Successfully")
    print("Saved :", filename)
    print("========================================")

    cursor.close()
    conn.close()


if __name__ == "__main__":
    generate_excel_report()